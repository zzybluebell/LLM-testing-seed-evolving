#!/usr/bin/env python3
"""Build a deliberately correct out/ from truth.json to self-test check.py (expects 13/13).

usage: make_reference_out.py [RUN_DIR]   (default tests/reference_run; writes RUN_DIR/work/out/)
"""
import json
import sys
from pathlib import Path

import matplotlib
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
T = json.loads((ROOT / "truth.json").read_text())
HEADER = ["month", "mrr", "new_customers", "churned_customers", "cogs", "sales_marketing_spend", "headcount"]


def rows():
    ws = load_workbook(ROOT / "data" / "financials.xlsx", read_only=True)["raw"]
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]


def build_xlsx(path, data):
    wb = Workbook()
    raw = wb.active
    raw.title = "raw"
    raw.append(HEADER)
    for r in data:
        raw.append(r)
    n, last = len(data), len(data) + 1
    ue = wb.create_sheet("unit_economics")
    ue.append(["Inputs"]), ue.append(["Starting active customers", 120]), ue.append([])
    ue.append(["month", "opening", "active", "gross_margin", "churn"])
    for i in range(n):
        r, src = i + 5, i + 2
        ue.append([f"=raw!A{src}", "=B2" if i == 0 else f"=C{r - 1}", f"=B{r}+raw!C{src}-raw!D{src}",
                   f"=(raw!B{src}-raw!E{src})/raw!B{src}", f"=raw!D{src}/B{r}"])
    end, b = n + 4, n + 6
    ue.append([])
    for label, f in [("CAC (TTM)", f"=SUM(raw!F{last - 11}:F{last})/SUM(raw!C{last - 11}:C{last})"),
                     ("ARPA", f"=raw!B{last}/C{end}"), ("Avg gross margin (12m)", f"=AVERAGE(D{end - 11}:D{end})"),
                     ("Avg churn (12m)", f"=AVERAGE(E{end - 11}:E{end})"), ("LTV", f"=B{b + 1}*B{b + 2}/B{b + 3}"),
                     ("LTV/CAC", f"=B{b + 4}/B{b}"), ("CAC payback months", f"=B{b}/(B{b + 1}*B{b + 2})")]:
        ue.append([label, f])
    d = wb.create_sheet("dcf")
    for row in (["Assumptions"], ["Revenue year 0", f"=SUM(raw!B{last - 11}:B{last})"], ["Discount rate", 0.12],
                ["Initial investment", -5000000], ["Year", 1, 2, 3, 4, 5], ["Growth", .25, .21, .17, .13, .10],
                ["FCF margin", .15, .175, .20, .225, .25], [],
                ["Revenue", "=B2*(1+B6)", "=B9*(1+C6)", "=C9*(1+D6)", "=D9*(1+E6)", "=E9*(1+F6)"],
                ["FCF", "=B9*B7", "=C9*C7", "=D9*D7", "=E9*E7", "=F9*F7"],
                ["Cash flows", "=B4", "=B10", "=C10", "=D10", "=E10", "=F10"],
                ["NPV", "=B4+NPV(B3,B10:F10)"], ["IRR", "=IRR(B11:G11)"]):
        d.append(row)
    lo = wb.create_sheet("loan")
    for row in (["Assumptions"], ["Principal", 2000000], ["Annual rate", 0.07], ["Payments", 60],
                ["Monthly payment (PMT)", "=PMT(B3/12,B4,-B2)"], [], ["Period", "Payment", "Interest", "Principal", "Balance"],
                [0, None, None, None, "=B2"]):
        lo.append(row)
    for p in range(1, 61):
        r = p + 8
        lo.append([p, "=$B$5", f"=E{r - 1}*$B$3/12", f"=B{r}-C{r}", f"=E{r - 1}-D{r}"])
    lo.append([]), lo.append(["Total interest", "=SUM(C9:C68)"])
    wb.save(path)


def build_charts(out, data):
    months = [r[0] for r in data]
    for name, series, labels in (("mrr", [[r[1] for r in data]], ["MRR"]),
                                 ("customers", [[r[2] for r in data], [r[3] for r in data]], ["New", "Churned"])):
        fig, ax = plt.subplots(figsize=(8, 4))
        for s, lab in zip(series, labels):
            ax.plot(months, s, label=lab)
        ax.set_xlabel("month"), ax.set_ylabel(name), ax.legend()
        ax.set_xticks(months[::6])
        fig.savefig(out / "charts" / f"{name}.png", dpi=100), plt.close(fig)


def add_text(slide, title, lines):
    slide.shapes.title.text = title
    box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(8.5), Inches(4)).text_frame
    box.text = lines[0]
    for line in lines[1:]:
        box.add_paragraph().text = line


def build_deck(out):
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text, title.placeholders[1].text = "Bluebell SaaS - Investor Update", "reference deck"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "MRR trend"
    s.shapes.add_picture(str(out / "charts" / "mrr.png"), Inches(1), Inches(1.5), width=Inches(8))
    add_text(prs.slides.add_slide(prs.slide_layouts[5]), "Gross margin & cost structure",
             [f"Average gross margin (36 months): {T['mean_gross_margin_36'] * 100:.1f}%",
              f"Revenue year 0: ${T['revenue_y0']:,.2f}"])
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Customer growth"
    s.shapes.add_picture(str(out / "charts" / "customers.png"), Inches(1), Inches(1.5), width=Inches(8))
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Unit economics"
    tbl = s.shapes.add_table(6, 2, Inches(1), Inches(1.6), Inches(7), Inches(3)).table
    for i, (k, v) in enumerate([("Metric", "Value"), ("CAC", f"${T['CAC_ttm']:,.2f}"), ("LTV", f"${T['LTV']:,.2f}"),
                                ("LTV/CAC", f"{T['LTV_over_CAC']:.2f}x"), ("CAC payback months", f"{T['payback_months']:.2f}"),
                                ("ARPA", f"${T['ARPA']:,.2f}")]):
        tbl.cell(i, 0).text, tbl.cell(i, 1).text = k, v
    add_text(prs.slides.add_slide(prs.slide_layouts[5]), "DCF",
             [f"NPV: ${T['NPV']:,.2f}", f"IRR: {T['IRR'] * 100:.2f}%"])
    add_text(prs.slides.add_slide(prs.slide_layouts[5]), "Loan",
             [f"Monthly payment: ${T['loan']['PMT']:,.2f}", f"Total interest: ${T['loan']['total_interest']:,.2f}"])
    add_text(prs.slides.add_slide(prs.slide_layouts[5]), "Use of funds", ["60% product", "30% go-to-market", "10% ops"])
    add_text(prs.slides.add_slide(prs.slide_layouts[5]), "Appendix: assumptions",
             ["Starting active customers: 120", "Discount rate 12%, no terminal value", "Loan 2,000,000 at 7% over 60 months",
              f"Conflict with last quarter's board deck (slide 7): it shows CAC ${T['slide7_cac_shown']:,}; "
              f"corrected value from this model is ${T['CAC_ttm']:,.2f}."])
    prs.save(out / "investor_update.pptx")


def main():
    run_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "tests" / "reference_run"
    out = run_dir / "work" / "out"
    (out / "charts").mkdir(parents=True, exist_ok=True)
    data = rows()
    build_xlsx(out / "model.xlsx", data), build_charts(out, data), build_deck(out)
    print("reference out/ written to", out)


if __name__ == "__main__":
    main()
