"""Step 2: build out/model.xlsx with live Excel formulas (no pasted metric values).

Sheets: raw, unit_economics, dcf, loan. Every computed metric is a formula
referencing the raw sheet (or other formula cells on the same sheet).
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]

NAVY = "1F2A44"
LIGHT = "E9ECF3"
ACCENT = "2E5EAA"
WHITE = "FFFFFF"

thin = Side(style="thin", color="B9C0CE")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
title_font = Font(name="Calibri", size=14, bold=True, color=NAVY)
hdr_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
bold = Font(name="Calibri", size=11, bold=True)
reg = Font(name="Calibri", size=11)
hdr_fill = PatternFill("solid", fgColor=NAVY)
sub_fill = PatternFill("solid", fgColor=LIGHT)
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

CUR = '"$"#,##0.00'
CUR0 = '"$"#,##0'
PCT = '0.00%'
NUM2 = '#,##0.00'

wb = openpyxl.Workbook()

# ---------------- raw ----------------
src = openpyxl.load_workbook(ROOT / "data" / "financials.xlsx", data_only=True)
src_raw = src["raw"]
raw = wb.active
raw.title = "raw"
for r_i, row in enumerate(src_raw.iter_rows(values_only=True), start=1):
    for c_i, val in enumerate(row, start=1):
        cell = raw.cell(row=r_i, column=c_i, value=val)
        if r_i == 1:
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
        else:
            cell.font = reg
            if c_i == 1:
                cell.alignment = left
            else:
                cell.alignment = Alignment(horizontal="right")
                if c_i in (2, 5, 6):
                    cell.number_format = CUR
for c, w in zip("ABCDEFG", (10, 13, 15, 18, 12, 22, 11)):
    raw.column_dimensions[c].width = w
raw.freeze_panes = "A2"

# rows: header=1, data 2..37 ; TTM window = rows 26..37 (2025-10..2026-09)

# ---------------- unit_economics ----------------
ue = wb.create_sheet("unit_economics")
ue["A1"] = "Bluebell SaaS Pte Ltd - Unit Economics"
ue["A1"].font = title_font
ue["A2"] = "Starting active customers (pre-period, 2023-09-30)"
ue["A2"].font = bold
ue["B2"] = 120
ue["B2"].font = bold
ue["A3"] = "Trailing-12-month window"
ue["A3"].font = bold
ue["B3"] = "2025-10 to 2026-09 (raw!26:37)"
ue["B3"].font = reg

headers = ["month", "active_customers", "gross_margin", "monthly_churn"]
for j, h in enumerate(headers, start=1):
    c = ue.cell(row=5, column=j, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = box

for i in range(36):
    g = 6 + i          # grid row
    r = 2 + i          # raw row
    ue.cell(row=g, column=1, value=f"=raw!A{r}")
    if i == 0:
        ue.cell(row=g, column=2, value=f"=$B$2+raw!C{r}-raw!D{r}")
        ue.cell(row=g, column=4, value=f"=raw!D{r}/$B$2")
    else:
        ue.cell(row=g, column=2, value=f"=B{g-1}+raw!C{r}-raw!D{r}")
        ue.cell(row=g, column=4, value=f"=raw!D{r}/B{g-1}")
    ue.cell(row=g, column=3, value=f"=(raw!B{r}-raw!E{r})/raw!B{r}")
    ue.cell(row=g, column=2).number_format = "#,##0"
    ue.cell(row=g, column=3).number_format = PCT
    ue.cell(row=g, column=4).number_format = PCT
    for j in range(1, 5):
        ue.cell(row=g, column=j).border = box
        ue.cell(row=g, column=j).font = reg

# summary block (F:G)
F, G = 6, 7
ue.cell(row=4, column=F, value="Metric").font = hdr_font
ue.cell(row=4, column=G, value="Value").font = hdr_font
for col in (F, G):
    cc = ue.cell(row=4, column=col)
    cc.fill = hdr_fill
    cc.alignment = center
    cc.border = box

summary = [
    ("S&M spend, trailing 12m",          "=SUM(raw!F26:F37)",                   CUR),
    ("New customers, trailing 12m",      "=SUM(raw!C26:C37)",                   "#,##0"),
    ("CAC (trailing 12 months)",         "=G5/G6",                              CUR),
    ("Average gross margin (L12)",       "=AVERAGE(C30:C41)",                   PCT),
    ("Average monthly churn (L12)",      "=AVERAGE(D30:D41)",                   PCT),
    ("Last month MRR (2026-09)",         "=raw!B37",                            CUR),
    ("Last month active customers",      "=B41",                                "#,##0"),
    ("ARPA (last MRR / last active)",    "=G10/G11",                            CUR),
    ("LTV = ARPA x avg GM / avg churn",  "=G12*G8/G9",                          CUR),
    ("LTV / CAC",                        "=G13/G7",                             NUM2),
    ("CAC payback (months)",             "=G7/(G12*G8)",                        NUM2),
]
for k, (label, formula, fmt) in enumerate(summary):
    rr = 5 + k
    a = ue.cell(row=rr, column=F, value=label)
    b = ue.cell(row=rr, column=G, value=formula)
    a.font = reg
    b.font = bold
    b.number_format = fmt
    a.border = box
    b.border = box
    b.alignment = Alignment(horizontal="right")
    if k % 2 == 1:
        a.fill = sub_fill
        b.fill = sub_fill

ue.cell(row=18, column=F, value="Note: monthly churn = churned / prior month active customers (first month base = 120).").font = Font(italic=True, size=9, color="555555")
for c, w in zip("ABCDEFG", (40, 18, 15, 15, 3, 38, 16)):
    ue.column_dimensions[c].width = w
ue.freeze_panes = "A6"

# ---------------- dcf ----------------
dcf = wb.create_sheet("dcf")
dcf["A1"] = "Bluebell SaaS Pte Ltd - 5-Year DCF (no terminal value)"
dcf["A1"].font = title_font

dcf["A3"] = "Discount rate"
dcf["B3"] = 0.12
dcf["A4"] = "Initial investment (year 0)"
dcf["B4"] = -5_000_000
dcf["A5"] = "Revenue year 0 (sum of last 12 months MRR)"
dcf["B5"] = "=SUM(raw!B26:B37)"
dcf["B3"].number_format = PCT
dcf["B4"].number_format = CUR0
dcf["B5"].number_format = CUR
for rr in (3, 4, 5):
    dcf.cell(row=rr, column=1).font = bold
    dcf.cell(row=rr, column=2).font = bold

# growth / margin assumption rows (years 1-5 in B..F)
dcf["A7"] = "Revenue growth"
dcf["A8"] = "FCF margin"
for j, (g, m) in enumerate(zip((0.25, 0.21, 0.17, 0.13, 0.10), (0.15, 0.175, 0.20, 0.225, 0.25))):
    cg = dcf.cell(row=7, column=2 + j, value=g)
    cm = dcf.cell(row=8, column=2 + j, value=m)
    cg.number_format = PCT
    cm.number_format = PCT
    cg.font = reg
    cm.font = reg
dcf["A7"].font = bold
dcf["A8"].font = bold
dcf["G7"] = "<- years 1-5"
dcf["G7"].font = Font(italic=True, size=9, color="555555")
dcf["G8"] = "<- years 1-5"
dcf["G8"].font = Font(italic=True, size=9, color="555555")

# model table: years 0..5 in columns B..G
dcf["A10"] = "Year"
dcf["A11"] = "Revenue"
dcf["A12"] = "FCF margin"
dcf["A13"] = "Free cash flow"
dcf["A14"] = "PV of FCF"
for rr in range(10, 15):
    dcf.cell(row=rr, column=1).font = bold
for j, yr in enumerate(range(0, 6)):
    col = get_column_letter(2 + j)
    yc = dcf.cell(row=10, column=2 + j, value=yr)
    yc.alignment = center
    yc.font = bold
    if j == 0:
        dcf.cell(row=11, column=2, value="=B5")
        dcf.cell(row=12, column=2, value=None)
        dcf.cell(row=13, column=2, value="=B4")
    else:
        prev = get_column_letter(1 + j)      # column letter of prior year
        gcol = get_column_letter(1 + j)     # growth for year j sits in row7 col B..F = col index 1+j
        dcf.cell(row=11, column=2 + j, value=f"={prev}11*(1+{gcol}7)")
        dcf.cell(row=12, column=2 + j, value=f"={gcol}8")
        dcf.cell(row=13, column=2 + j, value=f"={col}11*{col}12")
    dcf.cell(row=14, column=2 + j, value=f"={col}13/(1+$B$3)^{col}10")
    dcf.cell(row=11, column=2 + j).number_format = CUR0
    dcf.cell(row=13, column=2 + j).number_format = CUR0
    dcf.cell(row=14, column=2 + j).number_format = CUR0
    dcf.cell(row=12, column=2 + j).number_format = PCT
    for rr in range(10, 15):
        dcf.cell(row=rr, column=2 + j).border = box
        dcf.cell(row=rr, column=2 + j).font = reg
for rr in range(10, 15):
    dcf.cell(row=rr, column=1).border = box
for cc in range(2, 8):
    dcf.cell(row=10, column=cc).fill = sub_fill

dcf["A16"] = "NPV (year-0 investment + NPV of years 1-5 FCF)"
dcf["A16"].font = bold
dcf["B16"] = "=NPV(B3,C13:G13)+B13"
dcf["B16"].number_format = CUR0
dcf["B16"].font = bold
dcf["A17"] = "IRR"
dcf["A17"].font = bold
dcf["B17"] = "=IRR(B13:G13)"
dcf["B17"].number_format = PCT
dcf["B17"].font = bold
dcf["A18"] = "Memo: sum of discounted FCF"
dcf["B18"] = "=SUM(B14:G14)"
dcf["B18"].number_format = CUR0
dcf["A19"] = "Assumption: year-0 FCF is the initial investment only; no terminal value."
dcf["A19"].font = Font(italic=True, size=9, color="555555")
for c, w in zip("ABCDEFG", (48, 14, 14, 14, 14, 14, 12)):
    dcf.column_dimensions[c].width = w

# ---------------- loan ----------------
ln = wb.create_sheet("loan")
ln["A1"] = "Bluebell SaaS Pte Ltd - Term Loan Amortization"
ln["A1"].font = title_font
inputs = [
    ("Principal", 2_000_000, CUR0),
    ("Annual interest rate", 0.07, PCT),
    ("Monthly interest rate", "=B3/12", '0.0000%'),
    ("Term (months)", 60, "0"),
    ("Equal monthly payment", "=PMT(B4,B5,-B2)", CUR),
]
for i, (label, val, fmt) in enumerate(inputs):
    rr = 2 + i
    ln.cell(row=rr, column=1, value=label).font = bold
    c = ln.cell(row=rr, column=2, value=val)
    c.number_format = fmt
    c.font = bold

hdr_row = 9
for j, h in enumerate(["Month", "Beginning balance", "Payment", "Interest", "Principal repaid", "Ending balance"], start=1):
    c = ln.cell(row=hdr_row, column=j, value=h)
    c.font = hdr_font
    c.fill = hdr_fill
    c.alignment = center
    c.border = box

for i in range(60):
    rr = 10 + i
    ln.cell(row=rr, column=1, value=(1 if i == 0 else f"=A{rr-1}+1"))
    ln.cell(row=rr, column=2, value=("=B2" if i == 0 else f"=F{rr-1}"))
    ln.cell(row=rr, column=3, value="=$B$6")
    ln.cell(row=rr, column=4, value=f"=B{rr}*$B$4")
    ln.cell(row=rr, column=5, value=f"=C{rr}-D{rr}")
    ln.cell(row=rr, column=6, value=f"=B{rr}-E{rr}")
    for j in range(1, 7):
        cc = ln.cell(row=rr, column=j)
        cc.border = box
        cc.font = reg
        if j >= 2:
            cc.number_format = CUR
        else:
            cc.alignment = center

tot = 71
ln.cell(row=tot, column=1, value="Totals").font = bold
ln.cell(row=tot, column=3, value="=SUM(C10:C69)").number_format = CUR
ln.cell(row=tot, column=4, value="=SUM(D10:D69)").number_format = CUR
ln.cell(row=tot, column=5, value="=SUM(E10:E69)").number_format = CUR
for j in (3, 4, 5):
    ln.cell(row=tot, column=j).font = bold
ln["A73"] = "Total interest"
ln["A73"].font = bold
ln["B73"] = "=SUM(D10:D69)"
ln["B73"].number_format = CUR
ln["B73"].font = bold
ln["A74"] = "Check: ending balance after payment 60"
ln["A74"].font = bold
ln["B74"] = "=F69"
ln["B74"].number_format = CUR
ln["A75"] = "Check: total payments - total principal repaid"
ln["B75"] = "=C71-E71"
ln["B75"].number_format = CUR
for c, w in zip("ABCDEF", (38, 18, 14, 14, 18, 16)):
    ln.column_dimensions[c].width = w
ln.freeze_panes = "A10"

out = ROOT / "out" / "model.xlsx"
wb.save(out)
print("saved", out)
