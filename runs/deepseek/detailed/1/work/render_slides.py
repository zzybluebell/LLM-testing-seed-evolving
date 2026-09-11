import io, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 144  # px per inch -> 1920 x 1080
W_PX, H_PX = 1920, 1080
REG = '/home/user/Desktop/Work/LLM-testing/.venv/lib/python3.13/site-packages/matplotlib/mpl-data/fonts/ttf/DejaVuSans.ttf'
BOLD = '/home/user/Desktop/Work/LLM-testing/.venv/lib/python3.13/site-packages/matplotlib/mpl-data/fonts/ttf/DejaVuSans-Bold.ttf'

_font_cache = {}
def get_font(pt, bold):
    key = (round(pt * 2), bold)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(BOLD if bold else REG, round(pt * 2))
    return _font_cache[key]


def emu_px(v):
    return int(round(Emu(v).inches * SCALE))


def _color(rgb):
    return (rgb[0], rgb[1], rgb[2])


def _text_size(font, s):
    bbox = font.getbbox(s)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_runs(runs, max_w):
    """Wrap a paragraph's runs into lines of (word, size, bold, color, wpx, spacepx)."""
    tokens = []
    for r in runs:
        f = get_font(r['size'], r['bold'])
        for part in r['text'].split(' '):
            if part == '':
                continue
            tokens.append((part, r['size'], r['bold'], r['color'],
                           f.getlength(part), f.getlength(' ')))
    lines = []
    cur = []
    cur_w = 0
    for tok in tokens:
        wpx, spx = tok[4], tok[5]
        if cur and cur_w + spx + wpx > max_w:
            lines.append(cur)
            cur = [tok]
            cur_w = wpx
        else:
            if cur:
                cur_w += spx
            cur.append(tok)
            cur_w += wpx
    if cur:
        lines.append(cur)
    return lines


def draw_text(draw, text_frame, x, y, w, h):
    tf = text_frame
    anchor = tf.vertical_anchor
    lines = []
    line_heights = []
    for p in tf.paragraphs:
        runs = []
        for r in p.runs:
            size = r.font.size.pt if r.font.size else 14
            if r.font.color and r.font.color.type is not None:
                col = r.font.color.rgb
            else:
                col = (0x0B, 0x0B, 0x0B)
            runs.append({'text': r.text, 'size': size, 'bold': bool(r.font.bold),
                         'color': col})
        if not runs:
            runs.append({'text': ' ', 'size': 14, 'bold': False, 'color': (0, 0, 0)})
        wrapped = wrap_runs(runs, w)
        for ln in wrapped:
            lines.append((ln, p.alignment))
            line_heights.append(max(runs[0]['size'], 10) * 2 * 1.25)

    if not lines:
        return y

    total_h = sum(line_heights)
    if anchor is not None and str(anchor).endswith('MIDDLE'):
        cy = y + max(0, (h - total_h) / 2)
    elif anchor is not None and str(anchor).endswith('BOTTOM'):
        cy = y + max(0, h - total_h)
    else:
        cy = y

    for (ln, al), lh in zip(lines, line_heights):
        total_w = sum(t[4] for t in ln) + sum(t[5] for t in ln[:-1])
        if al is None:
            lx = x
        elif str(al).endswith('CENTER'):
            lx = x + (w - total_w) / 2
        elif str(al).endswith('RIGHT'):
            lx = x + w - total_w
        else:
            lx = x
        cx = lx
        for (word, s, b, col, wpx, spx) in ln:
            draw.text((cx, cy), word, font=get_font(s, b), fill=_color(col))
            cx += wpx + spx
        cy += lh
    return cy


def render_shape(draw, shape):
    l = emu_px(shape.left); t = emu_px(shape.top)
    w = emu_px(shape.width); h = emu_px(shape.height)
    st = shape.shape_type

    # picture
    if shape.shape_type == 13 or getattr(shape, 'image', None) is not None and hasattr(shape, '_element') and shape._element.tag.endswith('}pic'):
        try:
            img = Image.open(io.BytesIO(shape.image.blob)).convert('RGBA')
            img = img.resize((w, h), Image.LANCZOS)
            draw._image.paste(img, (l, t), img)
        except Exception as e:
            pass
        return
    if getattr(shape, 'has_text_frame', False) and not getattr(shape, 'has_table', False):
        # fill background if solid
        try:
            if shape.fill.type is not None and str(shape.fill.type) == 'MSO_FILL_TYPE.SOLID (1)':
                col = shape.fill.fore_color.rgb
                draw.rectangle([l, t, l + w, t + h], fill=_color(col))
        except Exception:
            pass
        draw_text(draw, shape.text_frame, l, t, w, h)
        return
    if getattr(shape, 'has_table', False):
        tbl = shape.table
        # cell fills + text
        col_ws = [emu_px(c.width) for c in tbl.columns]
        row_hs = [emu_px(r.height) for r in tbl.rows]
        cx = l
        for ci, cw in enumerate(col_ws):
            cy = t
            for ri, rh in enumerate(row_hs):
                cell = tbl.cell(ri, ci)
                # fill
                try:
                    if cell.fill.type is not None and str(cell.fill.type) == 'MSO_FILL_TYPE.SOLID (1)':
                        col = cell.fill.fore_color.rgb
                        draw.rectangle([cx, cy, cx + cw, cy + rh], fill=_color(col))
                except Exception:
                    pass
                # text
                cellx = cx + 12; celly = cy + 6
                draw_text(draw, cell.text_frame, cellx, celly, cw - 24, rh - 12)
                cy += rh
            cx += cw
        return
    # autoshape (accent bar, background rect) -> fill only
    try:
        if shape.fill.type is not None and str(shape.fill.type) == 'MSO_FILL_TYPE.SOLID (1)':
            col = shape.fill.fore_color.rgb
            draw.rectangle([l, t, l + w, t + h], fill=_color(col))
    except Exception:
        pass
    if getattr(shape, 'has_text_frame', False):
        draw_text(draw, shape.text_frame, l, t, w, h)


def render(pptx_path, out_dir):
    prs = Presentation(pptx_path)
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new('RGB', (W_PX, H_PX), (255, 255, 255))
        d = ImageDraw.Draw(img)
        d._image = img  # allow paste
        for shape in slide.shapes:
            render_shape(d, shape)
        out = os.path.join(out_dir, 'slide_%02d.png' % i)
        img.save(out)
        paths.append(out)
    return paths


if __name__ == '__main__':
    paths = render('out/investor_update.pptx', 'out/slides')
    print('rendered', len(paths), 'slides to out/slides/')
    for p in paths:
        print(' ', p, Image.open(p).size)
