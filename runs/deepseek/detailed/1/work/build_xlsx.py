import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = 'data/financials.xlsx'
OUT = 'out/model.xlsx'

# --- read source data ---
src = openpyxl.load_workbook(SRC, data_only=True)
src_ws = src['raw']
rows = list(src_ws.iter_rows(values_only=True))
header = rows[0]
data = rows[1:]
N = len(data)  # 36

wb = openpyxl.Workbook()

# ---- styles ----
bold = Font(bold=True)
title_font = Font(bold=True, size=14)
hdr_fill = PatternFill('solid', fgColor='D9E2F3')
asm_fill = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='C9C9C9')
border = Border(bottom=thin)

def style_header(cell):
    cell.font = bold
    cell.fill = hdr_fill
    cell.border = border
    cell.alignment = Alignment(horizontal='center')

# ============================================================
# Sheet 1: raw  (input data, unchanged)
# ============================================================
raw = wb.active
raw.title = 'raw'
for j, h in enumerate(header, start=1):
    c = raw.cell(row=1, column=j, value=h)
    style_header(c)
for i, row in enumerate(data, start=2):
    for j, v in enumerate(row, start=1):
        raw.cell(row=i, column=j, value=v)
# number formats for raw
for i in range(2, 2 + N):
    raw.cell(row=i, column=2).number_format = '#,##0.00'
    raw.cell(row=i, column=5).number_format = '#,##0.00'
    raw.cell(row=i, column=6).number_format = '#,##0.00'
for col, w in zip('ABCDEFG', (10, 12, 15, 17, 12, 21, 11)):
    raw.column_dimensions[col].width = w
raw.freeze_panes = 'A2'

# ============================================================
# Sheet 2: unit_economics  (live formulas referencing raw)
# ============================================================
ue = wb.create_sheet('unit_economics')
ue_hdr = ['month', 'active_customers', 'gross_margin', 'monthly_churn']
for j, h in enumerate(ue_hdr, start=1):
    c = ue.cell(row=1, column=j, value=h)
    style_header(c)

# time-series table rows 2..37
for i in range(2, 2 + N):
    r = i  # raw row == ue row (both data rows 2..37)
    ue.cell(row=i, column=1, value='=raw!A%d' % r)
    if i == 2:
        ue.cell(row=i, column=2, value='=120+raw!C2-raw!D2')
    else:
        ue.cell(row=i, column=2, value='=B%d+raw!C%d-raw!D%d' % (i - 1, r, r))
    ue.cell(row=i, column=3, value='=(raw!B%d-raw!E%d)/raw!B%d' % (r, r, r))
    if i == 2:
        ue.cell(row=i, column=4, value='=raw!D2/120')
    else:
        ue.cell(row=i, column=4, value='=raw!D%d/B%d' % (r, i - 1))
    ue.cell(row=i, column=3).number_format = '0.00%'
    ue.cell(row=i, column=4).number_format = '0.00%'
    ue.cell(row=i, column=2).number_format = '0'

# summary block (last-12-months = raw rows 26..37)
sr = 40
ue.cell(row=sr, column=1, value='Metric').font = bold
ue.cell(row=sr, column=2, value='Value').font = bold
ue.cell(row=sr, column=1).fill = hdr_fill
ue.cell(row=sr, column=2).fill = hdr_fill

summary = [
    ('CAC (trailing 12m)', '=SUM(raw!F26:F37)/SUM(raw!C26:C37)', '#,##0.00'),
    ('ARPA', '=raw!B37/(120+SUM(raw!C2:C37)-SUM(raw!D2:D37))', '#,##0.00'),
    ('Avg gross margin (last 12m)', '=AVERAGE(C26:C37)', '0.00%'),
    ('Avg monthly churn (last 12m)', '=AVERAGE(D26:D37)', '0.00%'),
    ('LTV', '=B42*B43/B44', '#,##0.00'),
    ('LTV / CAC', '=B45/B41', '0.00'),
    ('CAC payback (months)', '=B41/(B42*B43)', '0.00'),
]
for k, (label, formula, fmt) in enumerate(summary):
    rr = sr + 1 + k
    ue.cell(row=rr, column=1, value=label)
    c = ue.cell(row=rr, column=2, value=formula)
    c.number_format = fmt
    if label in ('LTV', 'LTV / CAC', 'CAC payback (months)'):
        c.font = bold

ue.column_dimensions['A'].width = 30
ue.column_dimensions['B'].width = 16
ue.column_dimensions['C'].width = 16
ue.column_dimensions['D'].width = 16
ue.freeze_panes = 'A2'

# ============================================================
# Sheet 3: dcf
# ============================================================
dcf = wb.create_sheet('dcf')
dcf.cell(row=1, column=1, value='DCF Model').font = title_font

dcf.cell(row=3, column=1, value='Assumptions').font = bold
dcf.cell(row=4, column=1, value='Revenue (year 0)')
dcf.cell(row=4, column=2, value='=SUM(raw!B26:B37)').number_format = '#,##0'
dcf.cell(row=5, column=1, value='Discount rate')
dcf.cell(row=5, column=2, value=0.12).number_format = '0%'
dcf.cell(row=6, column=1, value='Initial investment (year 0)')
dcf.cell(row=6, column=2, value=-5000000).number_format = '#,##0'

# growth / margin assumptions
hdr_row = 8
for j, h in enumerate(['Year', 'Growth', 'FCF margin'], start=1):
    c = dcf.cell(row=hdr_row, column=j, value=h)
    style_header(c)
growth = [0.25, 0.21, 0.17, 0.13, 0.10]
margins = [0.15, 0.175, 0.20, 0.225, 0.25]
for k in range(5):
    r = hdr_row + 1 + k
    dcf.cell(row=r, column=1, value='Year %d' % (k + 1))
    dcf.cell(row=r, column=2, value=growth[k]).number_format = '0%'
    dcf.cell(row=r, column=3, value=margins[k]).number_format = '0.0%'

# projection table
ph = 15
for j, h in enumerate(['Year', 'Revenue', 'FCF'], start=1):
    c = dcf.cell(row=ph, column=j, value=h)
    style_header(c)
dcf.cell(row=16, column=1, value='Year 0')
dcf.cell(row=16, column=2, value='=B4').number_format = '#,##0'
dcf.cell(row=16, column=3, value='=B6').number_format = '#,##0'
for k in range(5):
    r = 17 + k
    dcf.cell(row=r, column=1, value='Year %d' % (k + 1))
    dcf.cell(row=r, column=2, value='=B%d*(1+B%d)' % (r - 1, hdr_row + 1 + k)).number_format = '#,##0'
    dcf.cell(row=r, column=3, value='=B%d*C%d' % (r, hdr_row + 1 + k)).number_format = '#,##0'

dcf.cell(row=23, column=1, value='NPV').font = bold
dcf.cell(row=23, column=2, value='=C16+NPV(B5,C17:C21)').number_format = '#,##0'
dcf.cell(row=23, column=2).font = bold
dcf.cell(row=24, column=1, value='IRR').font = bold
dcf.cell(row=24, column=2, value='=IRR(C16:C21)').number_format = '0.00%'
dcf.cell(row=24, column=2).font = bold

dcf.column_dimensions['A'].width = 26
dcf.column_dimensions['B'].width = 16
dcf.column_dimensions['C'].width = 14

# ============================================================
# Sheet 4: loan
# ============================================================
ln = wb.create_sheet('loan')
ln.cell(row=1, column=1, value='Loan Amortization').font = title_font

ln.cell(row=3, column=1, value='Principal')
ln.cell(row=3, column=2, value=2000000).number_format = '#,##0'
ln.cell(row=4, column=1, value='Annual interest rate')
ln.cell(row=4, column=2, value=0.07).number_format = '0%'
ln.cell(row=5, column=1, value='Monthly interest rate')
ln.cell(row=5, column=2, value='=B4/12').number_format = '0.000000%'
ln.cell(row=6, column=1, value='Term (months)')
ln.cell(row=6, column=2, value=60)
ln.cell(row=7, column=1, value='Monthly payment').font = bold
ln.cell(row=7, column=2, value='=B3*B5/(1-(1+B5)^-B6)').number_format = '#,##0.00'
ln.cell(row=7, column=2).font = bold
ln.cell(row=8, column=1, value='Total interest').font = bold
ln.cell(row=8, column=2, value='=B7*B6-B3').number_format = '#,##0.00'
ln.cell(row=8, column=2).font = bold

hdr = ['Month', 'Beginning balance', 'Payment', 'Interest', 'Principal', 'Ending balance']
for j, h in enumerate(hdr, start=1):
    c = ln.cell(row=10, column=j, value=h)
    style_header(c)

for m in range(1, 61):
    r = 10 + m
    ln.cell(row=r, column=1, value=m)
    if m == 1:
        ln.cell(row=r, column=2, value='=$B$3')
    else:
        ln.cell(row=r, column=2, value='=F%d' % (r - 1))
    ln.cell(row=r, column=3, value='=$B$7')
    ln.cell(row=r, column=4, value='=B%d*$B$5' % r)
    ln.cell(row=r, column=5, value='=C%d-D%d' % (r, r))
    ln.cell(row=r, column=6, value='=B%d-E%d' % (r, r))
    for col in (2, 3, 4, 5, 6):
        ln.cell(row=r, column=col).number_format = '#,##0.00'

ln.column_dimensions['A'].width = 10
for col in 'BCDEF':
    ln.column_dimensions[col].width = 18
ln.freeze_panes = 'A11'

wb.save(OUT)
print('saved', OUT)
print('sheets:', wb.sheetnames)
