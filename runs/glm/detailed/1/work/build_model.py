import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

SRC = openpyxl.load_workbook('data/financials.xlsx', data_only=True)
src_ws = SRC['raw']
data = []
for r in src_ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None: continue
    data.append(r)
N = len(data)  # 36
print(f"{N} data rows")

wb = openpyxl.Workbook()

# ---------- styles ----------
H1 = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
H2 = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
HDR = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
BOLD = Font(name='Calibri', bold=True, size=10)
NORM = Font(name='Calibri', size=10)
NAVY = PatternFill('solid', fgColor='1F3864')
BLUE = PatternFill('solid', fgColor='2E5496')
LBLUE = PatternFill('solid', fgColor='D6E0F0')
LGREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BORD = Border(left=thin,right=thin,top=thin,bottom=thin)
C = Alignment(horizontal='center')
L = Alignment(horizontal='left')
R = Alignment(horizontal='right')

# ===================================================================
# Sheet 1: raw  (copy input data verbatim)
# ===================================================================
ws = wb.active
ws.title = 'raw'
headers = ['month','mrr','new_customers','churned_customers','cogs',
           'sales_marketing_spend','headcount']
for c,h in enumerate(headers,1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = HDR; cell.fill = BLUE; cell.alignment = C; cell.border = BORD
for i,row in enumerate(data,2):
    for c,val in enumerate(row,1):
        cell = ws.cell(row=i, column=c, value=val)
        cell.font = NORM; cell.border = BORD
        if c in (2,5,6): cell.number_format = '#,##0.00'
        if c in (3,4,7): cell.number_format = '#,##0'
        if c==1: cell.alignment = C
        else: cell.alignment = R
ws.column_dimensions['A'].width = 11
for col in 'BCDEFG':
    ws.column_dimensions[col].width = 16
ws.freeze_panes = 'A2'
FIRST = 2          # first data row
LAST  = N + 1      # last data row (37)

# ===================================================================
# Sheet 2: unit_economics  (live formulas referencing raw)
# ===================================================================
ue = wb.create_sheet('unit_economics')
ue['A1'] = 'Bluebell SaaS — Unit Economics (live formulas)'
ue['A1'].font = H1; ue['A1'].fill = NAVY
ue.merge_cells('A1:H1')
ue.row_dimensions[1].height = 24

# --- month-by-month computation block ---
cols = ['month','active_customers','gross_margin','monthly_churn',
        'CAC_trailing12','ARPA(monthly)','avg_gm_12','avg_churn_12']
hdr_row = 3
ue.cell(row=hdr_row, column=1, value='Monthly computation (formulas reference raw)')
ue.cell(row=hdr_row, column=1).font = BOLD
r0 = hdr_row + 1   # header row for table
for c,h in enumerate(cols,1):
    cell = ue.cell(row=r0, column=c, value=h)
    cell.font = HDR; cell.fill = BLUE; cell.alignment = C; cell.border = BORD
first = r0 + 1   # 5 = data row for month 1 (raw row 2)

for i in range(N):
    rr = first + i
    rawrow = FIRST + i   # raw sheet row
    # month
    ue.cell(row=rr, column=1, value=f"=raw!A{rawrow}").alignment = C
    # active_customers = prev active + new - churned ; month1 starts from 120
    if i == 0:
        ue.cell(row=rr, column=2, value=f"=120+raw!C{rawrow}-raw!D{rawrow}")
    else:
        ue.cell(row=rr, column=2, value=f"=B{rr-1}+raw!C{rawrow}-raw!D{rawrow}")
    # gross_margin = (mrr-cogs)/mrr
    ue.cell(row=rr, column=3, value=f"=(raw!B{rawrow}-raw!E{rawrow})/raw!B{rawrow}")
    # monthly_churn = churned / previous active
    if i == 0:
        ue.cell(row=rr, column=4, value=f"=raw!D{rawrow}/120")
    else:
        ue.cell(row=rr, column=4, value=f"=raw!D{rawrow}/B{rr-1}")
    # CAC trailing12 (blank for first 11)
    if i >= 11:
        rfrom = FIRST + i - 11
        rto = FIRST + i
        ue.cell(row=rr, column=5,
            value=f"=SUM(raw!F{rfrom}:F{rto})/SUM(raw!C{rfrom}:C{rto})")
    # ARPA monthly = mrr / active
    ue.cell(row=rr, column=6, value=f"=raw!B{rawrow}/B{rr}")
    # avg gross margin 12
    if i >= 11:
        ue.cell(row=rr, column=7, value=f"=AVERAGE(C{rr-11}:C{rr})")
    # avg churn 12
    if i >= 11:
        ue.cell(row=rr, column=8, value=f"=AVERAGE(D{rr-11}:D{rr})")
    for c in range(1,9):
        cell = ue.cell(row=rr, column=c)
        cell.font = NORM; cell.border = BORD
        if c==2: cell.number_format = '#,##0'
        if c in (3,4,7,8): cell.number_format = '0.0000'
        if c in (5,6): cell.number_format = '#,##0.00'
        if c>=2: cell.alignment = R

last_data = first + N - 1   # row for month 36

# --- headline unit-economics block (as of last month) ---
hr = last_data + 3
ue.cell(row=hr, column=1, value='Headline unit economics — as of latest month (2026-09)')
ue.cell(row=hr, column=1).font = H2; ue.cell(row=hr, column=1).fill = BLUE
ue.merge_cells(start_row=hr,start_column=1,end_row=hr,end_column=4)
LROW = last_data   # the last monthly row (month 36) used as the "as of" row

rows_def = [
    ('Active customers (latest)', f"=B{LROW}", '#,##0'),
    ('MRR (latest)', f"=raw!B{LAST}", '#,##0.00'),
    ('ARPA', f"=F{LROW}", '#,##0.00'),
    ('Average gross margin (trailing 12)', f"=G{LROW}", '0.00%'),
    ('Average monthly churn (trailing 12)', f"=H{LROW}", '0.00%'),
    ('CAC (trailing 12)', f"=E{LROW}", '#,##0.00'),
    ('LTV = ARPA x avgGM / avgChurn', f"=F{LROW}*G{LROW}/H{LROW}", '#,##0.00'),
    ('LTV / CAC', f"=(F{LROW}*G{LROW}/H{LROW})/E{LROW}", '0.00'),
    ('CAC payback (months) = CAC / (ARPA x avgGM)', f"=E{LROW}/(F{LROW}*G{LROW})", '0.00'),
]
for j,(label,formula,fmt) in enumerate(rows_def):
    rr = hr+1+j
    a = ue.cell(row=rr, column=1, value=label)
    a.font = BOLD; a.fill = LBLUE; a.border = BORD; a.alignment = L
    b = ue.cell(row=rr, column=2, value=formula)
    b.font = NORM; b.border = BORD; b.number_format = fmt; b.alignment = R
    # extend label across B? keep value in B; widen label col
ue.column_dimensions['A'].width = 38
for col in 'BCDEFGH':
    ue.column_dimensions[col].width = 15

# ===================================================================
# Sheet 3: dcf
# ===================================================================
dcf = wb.create_sheet('dcf')
dcf['A1'] = 'Bluebell SaaS — DCF / 5-Year FCF Model'
dcf['A1'].font = H1; dcf['A1'].fill = NAVY; dcf.merge_cells('A1:G1')
dcf.row_dimensions[1].height = 24

# assumptions block
dcf['A3'] = 'Assumptions'; dcf['A3'].font = BOLD
assum = [
    ('Revenue, Year 0 (sum last 12 mo MRR)', f"=SUM(raw!B{LAST-11}:B{LAST})", '#,##0'),
    ('Discount rate', 0.12, '0.0%'),
    ('Initial investment (Year 0)', -5000000, '#,##0'),
    ('Growth Y1', 0.25, '0.0%'),
    ('Growth Y2', 0.21, '0.0%'),
    ('Growth Y3', 0.17, '0.0%'),
    ('Growth Y4', 0.13, '0.0%'),
    ('Growth Y5', 0.10, '0.0%'),
    ('FCF margin Y1', 0.15, '0.0%'),
    ('FCF margin Y2', 0.175, '0.0%'),
    ('FCF margin Y3', 0.20, '0.0%'),
    ('FCF margin Y4', 0.225, '0.0%'),
    ('FCF margin Y5', 0.25, '0.0%'),
]
for j,(lab,val,fmt) in enumerate(assum):
    rr = 4 + j
    a = dcf.cell(row=rr, column=1, value=lab); a.font=NORM; a.alignment=L
    b = dcf.cell(row=rr, column=2, value=val); b.font=BOLD; b.number_format=fmt; b.alignment=R; b.fill=LBLUE
REV0_ROW = 4   # revenue year 0
DISC_ROW = 5
INV_ROW  = 6
GR_ROWS = [7,8,9,10,11]      # growth Y1..Y5
FM_ROWS = [12,13,14,15,16]   # fcf margin Y1..Y5

# FCF table
tr = 19
dcf.cell(row=tr, column=1, value='5-Year FCF').font = H2
dcf.cell(row=tr, column=1).fill = BLUE
dcf.merge_cells(start_row=tr,start_column=1,end_row=tr,end_column=7)
yr_hdr = tr+1
labels = ['Year','Revenue','Growth','FCF margin','FCF','Discount factor','PV of FCF']
for c,l in enumerate(labels,1):
    cell = dcf.cell(row=yr_hdr, column=c, value=l)
    cell.font=HDR; cell.fill=BLUE; cell.alignment=C; cell.border=BORD

# Year 0 row
y0 = yr_hdr+1
dcf.cell(row=y0, column=1, value='Year 0').alignment=C
dcf.cell(row=y0, column=2, value=f"=B{REV0_ROW}").number_format='#,##0'
dcf.cell(row=y0, column=3, value='—').alignment=C
dcf.cell(row=y0, column=4, value='—').alignment=C
dcf.cell(row=y0, column=5, value=f"=B{INV_ROW}").number_format='#,##0'  # FCF year0 = initial investment
dcf.cell(row=y0, column=6, value=1).number_format='0.0000'
dcf.cell(row=y0, column=7, value=f"=E{y0}*F{y0}").number_format='#,##0'

fcf_rows = []
for k in range(5):
    rr = y0+1+k
    prev = rr-1
    dcf.cell(row=rr, column=1, value=f'Year {k+1}').alignment=C
    # revenue = prev revenue * (1+growth)
    dcf.cell(row=rr, column=2, value=f"=B{prev}*(1+B{GR_ROWS[k]})").number_format='#,##0'
    dcf.cell(row=rr, column=3, value=f"=B{GR_ROWS[k]}").number_format='0.0%'
    dcf.cell(row=rr, column=4, value=f"=B{FM_ROWS[k]}").number_format='0.0%'
    dcf.cell(row=rr, column=5, value=f"=B{rr}*D{rr}").number_format='#,##0'
    dcf.cell(row=rr, column=6, value=f"=1/(1+B{DISC_ROW})^{k+1}").number_format='0.0000'
    dcf.cell(row=rr, column=7, value=f"=E{rr}*F{rr}").number_format='#,##0'
    fcf_rows.append(rr)

last_yr = y0+5+1-1  # last year row = y0+5
last_pv = last_yr
# NPV and IRR
nr = last_pv + 2
dcf.cell(row=nr, column=1, value='NPV (sum of PV of FCF incl. Year 0)').font=BOLD
dcf.cell(row=nr, column=1).fill=LBLUE
npv = dcf.cell(row=nr, column=5, value=f"=SUM(G{y0}:G{last_pv})")
npv.font=BOLD; npv.number_format='#,##0'; npv.fill=LBLUE
ir = nr+1
dcf.cell(row=ir, column=1, value='IRR').font=BOLD
dcf.cell(row=ir, column=1).fill=LBLUE
# IRR over the FCF cashflow stream (year0..year5)
irr = dcf.cell(row=ir, column=5, value=f"=IRR(E{y0}:E{last_pv})")
irr.font=BOLD; irr.number_format='0.0%'; irr.fill=LBLUE

dcf.column_dimensions['A'].width = 34
for col in 'BCDEFG':
    dcf.column_dimensions[col].width = 14

# ===================================================================
# Sheet 4: loan  (amortization)
# ===================================================================
ln = wb.create_sheet('loan')
ln['A1'] = 'Bluebell SaaS — Loan Amortization'
ln['A1'].font = H1; ln['A1'].fill = NAVY; ln.merge_cells('A1:F1')
ln.row_dimensions[1].height = 24
ln['A3']='Assumptions'; ln['A3'].font=BOLD
lass = [('Principal', 2000000, '#,##0'),
        ('Annual interest rate', 0.07, '0.00%'),
        ('Term (months)', 60, '#,##0')]
for j,(lab,val,fmt) in enumerate(lass):
    rr=4+j
    ln.cell(row=rr,column=1,value=lab).font=NORM
    b=ln.cell(row=rr,column=2,value=val); b.font=BOLD; b.number_format=fmt; b.fill=LBLUE
P_ROW=4; R_ROW=5; T_ROW=6
# monthly payment = PMT(rate/12, term, -principal)
mp_row=8
ln.cell(row=mp_row,column=1,value='Monthly payment (PMT)').font=BOLD
mp=ln.cell(row=mp_row,column=2,value=f"=PMT(B{R_ROW}/12,B{T_ROW},-B{P_ROW})")
mp.font=BOLD; mp.number_format='#,##0.00'; mp.fill=LBLUE

# table
thr=10
labs=['Month','Beginning balance','Payment','Interest','Principal','Ending balance']
for c,l in enumerate(labs,1):
    cell=ln.cell(row=thr,column=c,value=l); cell.font=HDR; cell.fill=BLUE; cell.alignment=C; cell.border=BORD
t0=thr+1
TERM=60
for m in range(1,TERM+1):
    rr=t0+m-1
    ln.cell(row=rr,column=1,value=m).alignment=C
    if m==1:
        ln.cell(row=rr,column=2,value=f"=B{P_ROW}")
    else:
        ln.cell(row=rr,column=2,value=f"=F{rr-1}")
    ln.cell(row=rr,column=3,value=f"=$B${mp_row}")
    ln.cell(row=rr,column=4,value=f"=B{rr}*B{R_ROW}/12")
    ln.cell(row=rr,column=5,value=f"=C{rr}-D{rr}")
    ln.cell(row=rr,column=6,value=f"=B{rr}-E{rr}")
    for c in range(1,7):
        cell=ln.cell(row=rr,column=c); cell.font=NORM; cell.border=BORD
        if c==1: cell.alignment=C
        else: cell.number_format='#,##0.00'; cell.alignment=R
last_t = t0+TERM-1
# total interest
ti=last_t+2
ln.cell(row=ti,column=1,value='Total interest').font=BOLD
tcell=ln.cell(row=ti,column=2,value=f"=SUM(D{t0}:D{last_t})")
tcell.font=BOLD; tcell.number_format='#,##0.00'; tcell.fill=LBLUE
ln.column_dimensions['A'].width=10
for col in 'BCDEF':
    ln.column_dimensions[col].width=16

wb.save('out/model.xlsx')
print('saved out/model.xlsx')
