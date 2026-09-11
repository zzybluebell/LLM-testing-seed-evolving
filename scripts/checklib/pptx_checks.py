"""python-pptx based inspection: pictures, tables, per-slide text."""
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PLACEHOLDER = re.compile(r"lorem|\bTBD\b|\bXXX\b|\[insert", re.I)


def open_deck(path):
    return Presentation(str(path))


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def is_picture(sh):
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    try:                                   # picture placeholders expose .image
        return sh.is_placeholder and getattr(sh, "image", None) is not None
    except Exception:  # noqa: BLE001 - placeholder without an image raises
        return False


def has_table(sh):
    return bool(getattr(sh, "has_table", False))


def slide_texts(prs):
    """One text blob per slide; table rows are joined with ' | ' so a label and its value stay adjacent."""
    texts = []
    for slide in prs.slides:
        lines = []
        for sh in iter_shapes(slide.shapes):
            if sh.has_text_frame:
                lines += [p.text for p in sh.text_frame.paragraphs if p.text.strip()]
            if has_table(sh):
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        texts.append("\n".join(lines))
    return texts


def deck_stats(prs):
    pics, max_rows = 0, 0
    for slide in prs.slides:
        for sh in iter_shapes(slide.shapes):
            pics += is_picture(sh)
            if has_table(sh):
                max_rows = max(max_rows, len(sh.table.rows))
    return {"slide_count": len(prs.slides), "picture_count": pics, "max_table_rows": max_rows}


def placeholder_hits(texts):
    return sorted({m.group(0) for t in texts for m in PLACEHOLDER.finditer(t)})
