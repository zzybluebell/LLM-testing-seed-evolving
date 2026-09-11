#!/usr/bin/env python3
"""Ground truth for the benchmark -> truth.json.

Definitions mirror prompts/detailed.md verbatim; see TASK.md section 1.2.
"""
import json
from pathlib import Path

import numpy_financial as npf
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "financials.xlsx"
OUT = ROOT / "truth.json"
ACTIVE_0 = 120
GROWTH = [0.25, 0.21, 0.17, 0.13, 0.10]
FCF_MARGIN = [0.15, 0.175, 0.20, 0.225, 0.25]
DISCOUNT_RATE, INITIAL_INVESTMENT = 0.12, -5_000_000
LOAN_PRINCIPAL, LOAN_RATE, LOAN_N = 2_000_000, 0.07, 60
SLIDE7_CAC_FACTOR = 1.35   # last quarter's board deck (slide 7) shows CAC at 1.35 x truth, whole dollars


def load_rows(path=DATA):
    ws = load_workbook(path, read_only=True)["raw"]
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    return [dict(zip(header, r)) for r in it if r and r[0] is not None]


def mean(xs):
    return sum(xs) / len(xs)


def compute(rows):
    active, opening, gm, churn = [], [], [], []
    prev = ACTIVE_0
    for r in rows:
        opening.append(prev)
        cur = prev + r["new_customers"] - r["churned_customers"]
        active.append(cur)
        gm.append((r["mrr"] - r["cogs"]) / r["mrr"])
        churn.append(r["churned_customers"] / prev)
        prev = cur
    last12 = rows[-12:]
    cac = sum(r["sales_marketing_spend"] for r in last12) / sum(r["new_customers"] for r in last12)
    arpa = rows[-1]["mrr"] / active[-1]
    gm12, churn12 = mean(gm[-12:]), mean(churn[-12:])
    ltv = arpa * gm12 / churn12
    payback = cac / (arpa * gm12)
    revenue_y0 = sum(r["mrr"] for r in last12)
    revenue, fcf, rev = [], [], revenue_y0
    for g, m in zip(GROWTH, FCF_MARGIN):
        rev *= 1 + g
        revenue.append(rev)
        fcf.append(rev * m)
    npv = INITIAL_INVESTMENT + sum(f / (1 + DISCOUNT_RATE) ** (y + 1) for y, f in enumerate(fcf))
    irr = float(npf.irr([INITIAL_INVESTMENT] + fcf))
    pmt = float(npf.pmt(LOAN_RATE / 12, LOAN_N, -LOAN_PRINCIPAL))
    return {
        "active_0": ACTIVE_0, "active": active, "opening_customers": opening,
        "active_last": active[-1], "gross_margin": gm, "churn": churn,
        "mean_gross_margin_36": mean(gm), "mean_gross_margin_12": gm12, "mean_churn_12": churn12,
        "CAC_ttm": cac, "ARPA": arpa, "LTV": ltv, "LTV_over_CAC": ltv / cac,
        "payback_months": payback,
        "revenue_y0": revenue_y0, "growth": GROWTH, "fcf_margin": FCF_MARGIN,
        "revenue": revenue, "fcf": fcf, "discount_rate": DISCOUNT_RATE,
        "initial_investment": INITIAL_INVESTMENT, "NPV": npv, "IRR": irr,
        "loan": {"principal": LOAN_PRINCIPAL, "annual_rate": LOAN_RATE, "n_payments": LOAN_N,
                 "PMT": pmt, "total_interest": LOAN_N * pmt - LOAN_PRINCIPAL},
        "slide7_cac_shown": round(cac * SLIDE7_CAC_FACTOR),
        "slide7_cac_factor": SLIDE7_CAC_FACTOR,
    }


def main():
    truth = compute(load_rows())
    OUT.write_text(json.dumps(truth, indent=2))
    keys = ("CAC_ttm", "ARPA", "LTV", "LTV_over_CAC", "payback_months", "NPV", "IRR",
            "mean_gross_margin_36", "slide7_cac_shown")
    print("wrote truth.json:", {k: round(truth[k], 4) for k in keys},
          "loan:", {k: round(v, 2) for k, v in truth["loan"].items()})


if __name__ == "__main__":
    main()
